from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).parent / '.dependencias'))

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'documentacao' / 'saida' / 'Proposta_Desenvolvimento_MOVEON.docx'
LOGO = ROOT / 'public' / 'logo.png'

RED = 'FE3000'
BLACK = '171717'
GRAY = '5F6368'
LIGHT = 'F4F5F7'
PALE_RED = 'FFF2EE'
WHITE = 'FFFFFF'
LINE = 'DADDE2'
USABLE_DXA = 9360


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_font(run, size=11, bold=False, color=BLACK, italic=False):
    run.font.name = 'Calibri'
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Calibri')
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=95, start=120, bottom=95, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for edge, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths)))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_ind = tbl_pr.find(qn('w:tblInd'))
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), '120')
    tbl_ind.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn('w:tcW'))
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn('w:w'), str(widths[idx]))
            tc_w.set(qn('w:type'), 'dxa')
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths, font_size=8.6):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, RED)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_font(run, font_size, True, WHITE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if row_index % 2:
                shade(cells[index], 'FAFAFB')
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(value))
            set_font(run, font_size, index == 0, BLACK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def para(doc, text='', bold_lead=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, 11, True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r, 11)
    else:
        r = p.add_run(text)
        set_font(r, 11)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    set_font(p.add_run(text), 11)


def number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    set_font(p.add_run(text), 11)


def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    return p


def callout(doc, label, text, fill=PALE_RED):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.25
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); p_pr.append(shd)
    borders = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single'); left.set(qn('w:sz'), '18'); left.set(qn('w:space'), '7'); left.set(qn('w:color'), RED)
    borders.append(left); p_pr.append(borders)
    r = p.add_run(f'{label}  ')
    set_font(r, 10.5, True, RED)
    r = p.add_run(text)
    set_font(r, 10.5, False, BLACK)


def add_picture(doc, path, width, alt):
    shape = doc.add_picture(str(path), width=width)
    shape._inline.docPr.set('name', alt[:120])
    shape._inline.docPr.set('descr', alt)
    return shape


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t'); text.text = '1'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_font(run, 8.5, False, GRAY)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'; normal.font.size = Pt(11); normal.font.color.rgb = rgb(BLACK)
    normal.paragraph_format.space_after = Pt(8); normal.paragraph_format.line_spacing = 1.333
    for style_name, size, before, after in [('Heading 1', 16, 18, 10), ('Heading 2', 13, 12, 6), ('Heading 3', 12, 8, 4)]:
        style = doc.styles[style_name]
        style.font.name = 'Calibri'; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = rgb(RED)
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run('PROPOSTA TÉCNICA  |  PORTAL MOVE.ON'), 8.5, True, GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run('MOVE.ON  •  Documento confidencial  •  Página '), 8.5, False, GRAY)
    add_page_field(footer)


def cover(doc):
    doc.add_paragraph().paragraph_format.space_after = Pt(22)
    if LOGO.exists():
        add_picture(doc, LOGO, Inches(2.25), 'Logotipo MOVE.ON em vermelho.')
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28); p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run('PROPOSTA TÉCNICA'), 12, True, RED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run('Desenvolvimento do Portal de Conteúdo / Blog MOVE.ON'), 25, True, BLACK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(25)
    set_font(p.add_run('Planejamento, implementação, homologação e publicação'), 13, False, GRAY)

    add_table(doc, ['INFORMAÇÃO', 'DEFINIÇÃO'], [
        ('Contratante', 'A definir na aprovação da proposta'),
        ('Proponente / responsável', 'A definir na formalização'),
        ('Início previsto', '13 de agosto de 2026'),
        ('Entrega final', 'Até 26 de setembro de 2026'),
        ('Prazo máximo', '45 dias corridos'),
        ('Dedicação planejada', 'Até 20 horas semanais; limite estimado de 120 horas'),
        ('Validade da proposta', 'Sujeita à aprovação formal e ao recebimento dos acessos e materiais necessários')
    ], [2600, 6760], 9.2)
    callout(doc, 'OBJETIVO', 'Entregar uma plataforma editorial moderna, segura, responsiva e administrável, integrada ao PostgreSQL e preparada para publicação, busca, métricas e relacionamento por newsletter.')
    doc.add_page_break()


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)
    cover(doc)

    heading(doc, '1. Apresentação da proposta', 1)
    para(doc, 'Esta proposta apresenta a abordagem para desenvolvimento e entrega do Portal de Conteúdo / Blog MOVE.ON. A solução reunirá o portal público, a área administrativa e o backend em uma aplicação integrada, com persistência PostgreSQL, organização modular, validações no servidor e mecanismos de segurança adequados ao tratamento de sessões, conteúdo e arquivos.')
    para(doc, 'O planejamento considera um MVP funcional e pronto para implantação, sem comprometer a evolução futura. A execução será conduzida em seis etapas verificáveis, cada uma com resultado esperado e ponto de validação pela contratante.')
    callout(doc, 'RESUMO EXECUTIVO', 'Prazo de até 45 dias corridos, de 13/08/2026 a 26/09/2026, com dedicação de até 20 horas semanais e limite estimado de até 120 horas.')

    heading(doc, '2. Objetivos', 1)
    bullet(doc, 'Disponibilizar uma página inicial com destaques, publicações recentes, categorias, pesquisa e navegação responsiva.')
    bullet(doc, 'Permitir leitura confortável, compartilhamento social, SEO por publicação e carregamento incremental do conteúdo.')
    bullet(doc, 'Oferecer um dashboard seguro para administrar publicações, categorias, aparência, administrador, newsletter e métricas reais.')
    bullet(doc, 'Armazenar dados no PostgreSQL por migrations e seeders reproduzíveis, mantendo organização profissional do código.')
    bullet(doc, 'Preparar a aplicação para implantação em Ubuntu Server com Docker Compose e configuração centralizada por variáveis de ambiente.')

    heading(doc, '3. Escopo funcional incluído', 1)
    add_table(doc, ['MÓDULO', 'ENTREGAS PRINCIPAIS'], [
        ('Portal público', 'Página inicial, destaques, cards padronizados, categorias, tema claro/escuro, busca, scroll infinito, retorno ao topo e responsividade.'),
        ('Publicações', 'Página individual, slug único, múltiplas categorias, data/hora de Brasília, imagens de capa, pré-visualização, rascunho, agendamento e destaque.'),
        ('Conteúdo rico', 'Formatação textual, fontes, tamanhos, cores, links, código, imagens e vídeos incorporados e redimensionáveis.'),
        ('Administração', 'Login, sessão segura, CRUD de categorias e publicações, identidade visual, perfil, foto, alteração protegida de e-mail e senha e logout.'),
        ('Busca e SEO', 'Full-text search PostgreSQL, URLs amigáveis, canonical, metadados essenciais, Open Graph, JSON-LD, sitemap e robots.'),
        ('Métricas', 'Visualizações, buscas, recortes por período, publicações mais acessadas e filtros para reduzir falsos positivos.'),
        ('Newsletter', 'Inscrição, modelo de e-mail editável, fila de envio, notificação após publicação e cancelamento por link seguro.'),
        ('Infraestrutura', 'Docker Compose com PostgreSQL, migrations, seed inicial, .env, documentação e roteiro de instalação em Ubuntu Server.')
    ], [2200, 7160], 8.6)

    heading(doc, '4. Abordagem técnica e decisões', 1)
    para(doc, 'A aplicação adotará uma arquitetura integrada de frontend e backend no mesmo repositório, com módulos separados por responsabilidade. Essa decisão simplifica implantação, versionamento, configuração e manutenção do MVP, sem impedir a separação física dos serviços caso o volume ou a operação futura exijam escalabilidade independente.')
    add_table(doc, ['DECISÃO', 'JUSTIFICATIVA'], [
        ('PostgreSQL', 'Integridade relacional, transações, índices avançados, JSONB e full-text search nativo.'),
        ('REST como padrão', 'Contratos simples, previsíveis e suficientes para o portal e o dashboard; evita complexidade operacional sem benefício imediato.'),
        ('Full-text search', 'Substitui LIKE/ILIKE por busca linguística indexada, com maior relevância e menor custo em volumes elevados.'),
        ('Paginação por cursor', 'Mantém custo mais estável no scroll infinito e evita descarte progressivo de linhas típico de OFFSET.'),
        ('Fila persistente de newsletter', 'Separa a publicação do envio de e-mail, permite novas tentativas e reduz o tempo da requisição administrativa.'),
        ('Otimização de mídia', 'Recompressão, formatos adequados, tamanhos controlados, carregamento lazy e imagem social 1200 × 630.'),
        ('Validação no backend', 'O servidor valida entradas, sessão, permissões e arquivos; o frontend não é tratado como fonte confiável.')
    ], [2300, 7060], 8.5)

    heading(doc, '5. Entregáveis', 1)
    number(doc, 'Código-fonte organizado do portal público, dashboard e backend.')
    number(doc, 'Migrations SQL, seed administrativo inicial e estrutura PostgreSQL.')
    number(doc, 'Diagramas de casos de uso, classes e modelagem do banco de dados.')
    number(doc, 'Configuração Docker Compose e variáveis documentadas em .env.example.')
    number(doc, 'Documentação técnica, instruções de execução e roteiro de instalação no servidor.')
    number(doc, 'Aplicação testada, homologada com a contratante e implantada no ambiente disponibilizado.')

    heading(doc, '6. Cronograma de execução', 1)
    para(doc, 'Prazo estimado de até 45 dias corridos. Considerando a aprovação da proposta e o recebimento dos acessos e materiais necessários, o início está previsto para 13 de agosto de 2026 e a entrega final para até 26 de setembro de 2026.')
    schedule = [
        ('1', '13 a 19/08', 'Modelo e validação', 'Protótipo, requisitos e escopo do MVP validados e aprovados pela contratante.', '20 h'),
        ('2', '20 a 26/08', 'Modelagem técnica', 'Diagrama de casos de uso, diagrama de classes, modelagem do banco de dados e definição das tabelas SQL.', '20 h'),
        ('3', '27/08 a 02/09', 'Estrutura e base da aplicação', 'Estrutura inicial do projeto, configurações, conexão com o banco de dados, autenticação administrativa e base da API.', '20 h'),
        ('4', '03 a 09/09', 'Painel administrativo e portal público', 'Gestão de categorias e publicações, listagem pública, página de artigo, filtros, busca e estrutura responsiva.', '20 h'),
        ('5', '10 a 16/09', 'Métricas e indicadores', 'Registro de visualizações e buscas, consolidação dos dados, indicadores e gráficos no painel administrativo.', '20 h'),
        ('6', '17 a 26/09', 'Homologação, ajustes e publicação', 'Testes, correções, ajustes previstos, validação da contratante, implantação e entrega final.', '20 h')
    ]
    add_table(doc, ['ETAPA', 'PERÍODO', 'ATIVIDADE', 'RESULTADO ESPERADO', 'HORAS'], schedule, [650, 1300, 2100, 4610, 700], 7.9)
    callout(doc, 'MARCOS PREVISTOS', 'Início em 13/08/2026 e entrega final até 26/09/2026, completando 45 dias corridos. A soma planejada é de 120 horas.')
    para(doc, 'As datas poderão ser reprogramadas proporcionalmente caso ocorram atrasos na aprovação, no fornecimento de acessos e materiais ou nas validações de responsabilidade da contratante.', bold_lead='As datas poderão ser reprogramadas')

    heading(doc, '7. Método de execução e acompanhamento', 1)
    add_table(doc, ['PASSO', 'COMO SERÁ EXECUTADO'], [
        ('1. Planejamento', 'Detalhamento do lote da semana, dependências e critérios de aceite.'),
        ('2. Implementação', 'Desenvolvimento incremental com validação técnica contínua.'),
        ('3. Demonstração', 'Apresentação do resultado da etapa para conferência da contratante.'),
        ('4. Validação', 'Registro de ajustes aderentes ao escopo e aprovação para avançar.'),
        ('5. Entrega', 'Homologação final, implantação, documentação e transferência dos artefatos.')
    ], [2200, 7160], 8.8)
    para(doc, 'O acompanhamento deverá ocorrer por um canal oficial definido entre as partes. Decisões, aprovações e alterações de escopo devem ser registradas por escrito para preservar prazo, rastreabilidade e entendimento comum.')

    heading(doc, '8. Responsabilidades das partes', 1)
    heading(doc, '8.1 Responsabilidades da proponente', 2)
    bullet(doc, 'Executar as atividades técnicas previstas no cronograma e comunicar impedimentos relevantes.')
    bullet(doc, 'Aplicar práticas de segurança, validação, organização do código, versionamento e documentação.')
    bullet(doc, 'Disponibilizar versões para validação e corrigir não conformidades relacionadas ao escopo aprovado.')
    bullet(doc, 'Preservar a confidencialidade dos acessos e materiais recebidos para execução do projeto.')
    heading(doc, '8.2 Responsabilidades da contratante', 2)
    bullet(doc, 'Aprovar a proposta, o escopo e os resultados intermediários nos prazos acordados.')
    bullet(doc, 'Fornecer logotipo, favicon, textos, imagens, credenciais, domínio, servidor e demais acessos necessários.')
    bullet(doc, 'Indicar uma pessoa responsável por consolidar decisões e devolutivas.')
    bullet(doc, 'Validar conteúdo, identidade, regras de negócio e implantação antes do aceite final.')

    heading(doc, '9. Premissas, dependências e limites', 1)
    add_table(doc, ['TEMA', 'CONDIÇÃO CONSIDERADA'], [
        ('Aprovações', 'A contratante fornecerá devolutivas em tempo hábil; atrasos deslocam as datas proporcionalmente.'),
        ('Infraestrutura', 'Servidor Ubuntu, domínio, DNS e credenciais serão fornecidos com permissão suficiente para implantação.'),
        ('Serviço de e-mail', 'Conta SMTP válida, reputação do domínio e limites do provedor são de responsabilidade da contratante.'),
        ('Conteúdo', 'Textos, imagens, direitos de uso e validação editorial serão fornecidos ou aprovados pela contratante.'),
        ('Escopo', 'Solicitações fora das entregas descritas serão avaliadas separadamente quanto a prazo e esforço.'),
        ('Terceiros', 'Custos de hospedagem, domínio, e-mail, CDN, armazenamento ou serviços externos não estão definidos nesta proposta.'),
        ('Compatibilidade', 'A validação considera versões atuais dos principais navegadores em computadores, tablets e celulares.')
    ], [2100, 7260], 8.5)

    heading(doc, '10. Segurança, privacidade e qualidade', 1)
    bullet(doc, 'Cookies de sessão HttpOnly, Secure em produção e SameSite configurável.')
    bullet(doc, 'Senhas protegidas por hash forte; alteração de senha revoga sessões existentes.')
    bullet(doc, 'Consultas parametrizadas, validação estrita das entradas e sanitização do conteúdo rico.')
    bullet(doc, 'Uploads validados por conteúdo real, recomprimidos e armazenados com nomes não previsíveis.')
    bullet(doc, 'Métricas filtram acessos administrativos e robôs identificados, reduzindo falsos positivos.')
    bullet(doc, 'Testes funcionais, de integração, responsividade, segurança e implantação compatíveis com o risco do MVP.')

    heading(doc, '11. Homologação e critérios de aceite', 1)
    para(doc, 'A entrega será considerada apta para aceite quando os itens abaixo estiverem demonstrados no ambiente de homologação ou produção disponibilizado:')
    bullet(doc, 'Portal público, busca, categorias, artigos, compartilhamento, SEO e navegação responsiva operando conforme o escopo.')
    bullet(doc, 'Dashboard autenticado com gestão de conteúdo, identidade, administrador, newsletter e métricas.')
    bullet(doc, 'Persistência PostgreSQL criada pelas migrations e dados iniciais carregados pelo seeder.')
    bullet(doc, 'Upload e otimização de imagens, estados de publicação e agendamento validados.')
    bullet(doc, 'Documentação e instruções de instalação entregues junto ao código-fonte.')
    bullet(doc, 'Não conformidades do escopo registradas durante a homologação corrigidas antes do encerramento.')

    heading(doc, '12. Condições comerciais e controle de mudanças', 1)
    callout(doc, 'INVESTIMENTO', 'O valor, a forma de pagamento, a titularidade contratual e eventuais tributos deverão ser preenchidos e aprovados pelas partes antes da assinatura. Esta versão não presume valores não informados.')
    para(doc, 'Qualquer mudança que altere requisitos, integrações, volume de conteúdo, infraestrutura, identidade aprovada ou entregáveis será analisada quanto a impacto. A execução dependerá de autorização escrita e, quando necessário, de aditivo de prazo e esforço.')
    para(doc, 'O limite estimado de 120 horas representa o planejamento do escopo descrito. A priorização do MVP poderá ser usada para preservar o prazo quando surgirem dependências externas, desde que aprovada pela contratante.')

    heading(doc, '13. Vigência da proposta e aceite', 1)
    para(doc, 'A execução poderá começar em 13 de agosto de 2026 se a proposta estiver aprovada e os materiais, acessos e responsáveis pela validação estiverem disponíveis. Caso a aprovação ocorra posteriormente, o cronograma deverá ser recalculado preservando a sequência e a duração relativa das etapas.')
    para(doc, 'A assinatura abaixo registra concordância com o escopo, o cronograma, as premissas e o procedimento de controle de mudanças. Os campos comerciais pendentes devem ser preenchidos antes da formalização.')

    doc.add_paragraph().paragraph_format.space_after = Pt(24)
    add_table(doc, ['PELA CONTRATANTE', 'PELA PROPONENTE'], [
        ('Nome: __________________________________', 'Nome: __________________________________'),
        ('Cargo: __________________________________', 'Cargo: __________________________________'),
        ('Assinatura: ______________________________', 'Assinatura: ______________________________'),
        ('Data: ____ / ____ / ______', 'Data: ____ / ____ / ______')
    ], [4680, 4680], 9.5)
    callout(doc, 'PRÓXIMO PASSO', 'Preencher os dados das partes e as condições comerciais, aprovar formalmente a proposta e disponibilizar acessos e materiais para o início previsto.')

    settings = doc.settings._element
    update = settings.find(qn('w:updateFields'))
    if update is None:
        update = OxmlElement('w:updateFields'); settings.append(update)
    update.set(qn('w:val'), 'true')
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()
